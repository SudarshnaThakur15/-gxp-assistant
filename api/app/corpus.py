"""Load corpus sources through one adapter per source type, and validate them.

Adapters (all produce the same metadata fields):
  - Markdown with a YAML header           corpus/northbridge/{shared,client-a,client-b}/*.md
  - PDF with a sidecar *.meta.yaml        corpus/public/*.pdf
  - Word files described by a DMS export  corpus/northbridge/dms-export/ (CSV + mapping.yaml)

Our own authored sources must be clean: any problem is an error that blocks
ingestion. Records arriving from a client's DMS fail closed instead: a record
with a problem (e.g. no status) is rejected and logged, never guessed at.
"""

import csv
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import yaml
from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]
CORPUS = REPO_ROOT / "corpus"
DMS_EXPORT = CORPUS / "northbridge" / "dms-export"

REQUIRED = ["code", "title", "version", "status", "doc_type", "owner", "tenant", "confidentiality"]
REQUIRED_KEYS = ["effective_date", "client"]  # must be present, may be null
STATUSES = {"draft", "effective", "superseded", "withdrawn"}
DOC_TYPES = {"sop", "policy", "specification", "quality_agreement", "deviation", "capa", "guideline"}
CLIENTS = {None, "client-a", "client-b"}
FOLDER_SCOPE = {  # folder -> (tenant, client)
    "public": ("public", None),
    "northbridge/shared": ("northbridge", None),
    "northbridge/client-a": ("northbridge", "client-a"),
    "northbridge/client-b": ("northbridge", "client-b"),
}
BANNER = "SYNTHETIC – DEMO ONLY"
DOC_REF = re.compile(r"\b(?:SOP-[A-Z]{2,3}-\d{3}|DEV-\d{4}-\d{3}|CAPA-\d{4}-\d{3}|QA-[A-Z]{3}-\d{3}|"
                     r"SPEC-[A-Z]{3}-\d{3}|POL-[A-Z]{2}-\d{3})\b")


@dataclass
class Source:
    path: Path        # the content file: .md, .pdf or .docx
    format: str       # markdown | pdf | docx
    meta: dict
    problems: list[str] = field(default_factory=list)

    @property
    def key(self) -> tuple[str, str]:
        return self.meta.get("code"), str(self.meta.get("version"))

    @property
    def label(self) -> str:
        return str(self.path.relative_to(REPO_ROOT))

    @property
    def from_dms(self) -> bool:
        return self.format == "docx"


def body_text(source: Source) -> str:
    """Document text used for banner and cross-reference checks (not for PDFs)."""
    if source.format == "markdown":
        return source.path.read_text().split("---\n", 2)[2]
    if source.format == "docx":
        doc = Document(source.path)
        cells = (cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        return "\n".join([p.text for p in doc.paragraphs] + list(cells))
    return ""


# --- adapters ---------------------------------------------------------------

def load_markdown(path: Path) -> Source:
    text = path.read_text()
    if not text.startswith("---\n"):
        return Source(path, "markdown", {}, ["missing YAML header"])
    return Source(path, "markdown", yaml.safe_load(text.split("---\n", 2)[1]) or {})


def load_pdf(meta_path: Path) -> Source:
    pdf = meta_path.with_name(meta_path.name.removesuffix(".meta.yaml") + ".pdf")
    return Source(pdf, "pdf", yaml.safe_load(meta_path.read_text()) or {})


def load_dms_export(folder: Path) -> list[Source]:
    mapping = yaml.safe_load((folder / "mapping.yaml").read_text())
    values = {name: {k.casefold(): v for k, v in table.items()} for name, table in mapping["values"].items()}
    files_dir = (folder / "files").resolve()
    sources = []
    # utf-8-sig: Excel-saved CSVs start with a byte-order mark that would corrupt the first column name
    with (folder / "dms_export.csv").open(newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            meta, problems = {"tenant": mapping["tenant"]}, []
            for name, column in mapping["fields"].items():
                raw = (row.get(column) or "").strip()
                value = values[name].get(raw.casefold(), raw) if raw and name in values else (raw or None)
                if raw and name in values and raw.casefold() not in values[name]:
                    problems.append(f"{name} value {raw!r} has no mapping")
                if value and name in ("effective_date", "review_date"):
                    try:
                        value = date.fromisoformat(value)
                    except ValueError:
                        problems.append(f"{name} {raw!r} is not a date")
                meta[name] = value
            path = (files_dir / (row.get(mapping["file_column"]) or "")).resolve()
            if not path.is_relative_to(files_dir) or path == files_dir:  # e.g. "../../secret"
                problems.append("file name points outside the export folder")
                path = files_dir / "INVALID"
            sources.append(Source(path, "docx", meta, problems))
    return sources


def load_sources() -> list[Source]:
    sources = [load_markdown(p) for p in sorted((CORPUS / "northbridge").glob("*/*.md"))]
    sources += [load_pdf(p) for p in sorted((CORPUS / "public").glob("*.meta.yaml"))]
    if DMS_EXPORT.exists():
        sources += load_dms_export(DMS_EXPORT)
    return sources


# --- validation -------------------------------------------------------------

def check_source(source: Source) -> None:
    m, p = source.meta, source.problems
    p += [f"missing {k}" for k in REQUIRED if not m.get(k)]
    p += [f"missing {k}" for k in REQUIRED_KEYS if k not in m]
    if m.get("status") and m["status"] not in STATUSES:
        p.append(f"unknown status {m['status']!r}")
    if m.get("doc_type") and m["doc_type"] not in DOC_TYPES:
        p.append(f"unknown doc_type {m['doc_type']!r}")
    for k in ("effective_date", "review_date"):
        if m.get(k) is not None and not isinstance(m[k], date):
            p.append(f"{k} is not a date")
    if not source.path.exists():
        p.append(f"file not found: {source.path.name}")
        return

    if source.from_dms:
        if m.get("client") not in CLIENTS:
            p.append(f"unknown client {m.get('client')!r}")
    else:
        folder = str(source.path.parent.relative_to(CORPUS))
        if FOLDER_SCOPE.get(folder) != (m.get("tenant"), m.get("client")):
            p.append(f"folder {folder!r} does not match tenant/client {m.get('tenant')}/{m.get('client')}")

    if source.format == "markdown" and source.path.name != f"{m.get('code')}_v{m.get('version')}.md":
        p.append(f"file name should be {m.get('code')}_v{m.get('version')}.md")
    if source.format in ("markdown", "docx") and m.get("tenant") == "northbridge" and BANNER not in body_text(source):
        p.append("missing synthetic-data banner")
    if m.get("regulatory_status") == "consultation_draft" and not m.get("required_qualifier"):
        p.append("draft guidance needs a required_qualifier")


def check_corpus(accepted: list[Source], errors: list[str]) -> None:
    for key, n in Counter(s.key for s in accepted).items():
        if n > 1:
            errors.append(f"duplicate document {key[0]} v{key[1]}")

    by_code = defaultdict(list)
    for s in accepted:
        by_code[s.meta["code"]].append(s)
    for code, versions in by_code.items():
        effective = [s for s in versions if s.meta["status"] == "effective"]
        if len(effective) != 1:
            errors.append(f"{code}: expected exactly one effective version, found {len(effective)}")
        for s in versions:
            nxt = s.meta.get("superseded_by")
            if s.meta["status"] == "superseded" and not any(str(v.meta["version"]) == str(nxt) for v in versions):
                errors.append(f"{s.label}: superseded_by {nxt!r} does not exist")

    clients = {s.meta["code"]: s.meta["client"] for s in accepted}
    for s in accepted:
        for ref in set(DOC_REF.findall(body_text(s))) - {s.meta["code"]}:
            if ref not in clients:
                errors.append(f"{s.label}: references unknown document {ref}")
            elif s.meta["client"] and clients[ref] and clients[ref] != s.meta["client"]:
                errors.append(f"{s.label}: references {ref} of another client ({clients[ref]})")


def load_and_validate() -> tuple[list[Source], list[Source], list[str]]:
    """Return (accepted, rejected, errors). Errors block ingestion; rejections are DMS records that fail closed."""
    accepted, rejected, errors = [], [], []
    sources = load_sources()
    # Conflicting DMS records (same code+version twice, or two effective versions) are rejected, not guessed.
    dms = [s for s in sources if s.from_dms]
    for key, n in Counter(s.key for s in dms).items():
        if n > 1:
            for s in dms:
                if s.key == key:
                    s.problems.append("duplicate record in DMS export")
    for code, n in Counter(s.meta.get("code") for s in dms if s.meta.get("status") == "effective").items():
        if n > 1:
            for s in dms:
                if s.meta.get("code") == code and s.meta.get("status") == "effective":
                    s.problems.append("more than one effective version in DMS export")
    for source in sources:
        check_source(source)
        if not source.problems:
            accepted.append(source)
        elif source.from_dms:
            rejected.append(source)
        else:
            errors += [f"{source.label}: {problem}" for problem in source.problems]
    check_corpus(accepted, errors)
    return accepted, rejected, errors
